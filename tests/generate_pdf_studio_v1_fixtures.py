from pathlib import Path
import json
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from PIL import Image, ImageDraw
from docx import Document
from zipfile import ZipFile, ZIP_DEFLATED, ZIP_STORED

def make_pdf(path, count, label):
    c = canvas.Canvas(path, pagesize=A4)
    for i in range(1, count + 1):
        c.setFont("Helvetica", 18)
        c.drawString(72, 780, f"{label} — page {i}/{count}")
        c.setFont("Helvetica", 10)
        c.drawString(72, 752, "Synthetic nLab PDF Studio 1.0 acceptance fixture")
        c.showPage()
    c.save()

make_pdf("/tmp/pdf-studio-v1.pdf", 12, "Fixture A")
make_pdf("/tmp/pdf-studio-v1-b.pdf", 3, "Fixture B")

img = Image.new("RGBA", (500, 180), (255, 255, 255, 0))
draw = ImageDraw.Draw(img)
draw.line((30, 120, 130, 45, 250, 125, 430, 55), fill=(30, 65, 105, 255), width=8)
img.save("/tmp/pdf-studio-v1-signature.png")

Path("/tmp/pdf-studio-v1-config.json").write_text(json.dumps({
    "schema": "nlab-pdf-studio/v1",
    "variables": {"values": {"INITIALS": "CI", "DATE_A": "2026-09-27"}},
    "naming": {"prefix": "CI_", "template": "{FILENAME}", "suffix": "_TEST"}
}, indent=2))


doc = Document()
doc.add_heading("nLab PDF Studio 1.0 DOCX fixture", level=1)
doc.add_paragraph("Synthetic DOCX source converted to PDF in browser.")
doc.add_paragraph("Reference: DOCX-TO-PDF-CI")
doc.save("/tmp/pdf-studio-v1.docx")

with ZipFile("/tmp/pdf-studio-v1.odt", "w") as z:
    z.writestr("mimetype", "application/vnd.oasis.opendocument.text", compress_type=ZIP_STORED)
    z.writestr("META-INF/manifest.xml", """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">
<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>
<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>""", compress_type=ZIP_DEFLATED)
    z.writestr("content.xml", """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.3">
<office:body><office:text><text:h text:outline-level="1">nLab PDF Studio 1.0 ODT fixture</text:h><text:p>Synthetic ODT source converted to PDF in browser.</text:p><text:p>Reference: ODT-TO-PDF-CI</text:p></office:text></office:body>
</office:document-content>""", compress_type=ZIP_DEFLATED)
