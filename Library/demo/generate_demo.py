#!/usr/bin/env python3
"""Generate the public nLab synthetic demo corpus.
All records are intentionally synthetic. No real personal data is used.
"""
from pathlib import Path
import argparse, csv, hashlib, io, json, random, shutil, zipfile
from datetime import datetime
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, landscape
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

REG = ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"]
ITALIC = ["/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Italic.ttf"]
BOLD = ["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"]

def find_font(names, size):
    for p in names:
        if Path(p).exists():
            return ImageFont.truetype(p, size=size)
    return ImageFont.load_default()

def fnt(which, size):
    return find_font({"reg":REG,"italic":ITALIC,"bold":BOLD}[which], size)

def ensure_dirs(root):
    for d in [
        "demo-input/PDF", "demo-input/OCR", "demo-input/Images", "demo-input/Office",
        "demo-input/Data", "demo-input/Zip", "demo-input/Compare/tree-A/docs",
        "demo-input/Compare/tree-A/data", "demo-input/Compare/tree-B/docs",
        "demo-input/Compare/tree-B/data", "demo-output", "docs"]:
        (root/d).mkdir(parents=True, exist_ok=True)

def save_svg(path):
    svg = '<svg xmlns="http://www.w3.org/2000/svg" width="900" height="600" viewBox="0 0 900 600"><rect width="900" height="600" fill="#ffffff"/><rect x="30" y="30" width="840" height="540" rx="24" fill="#f5f7fa" stroke="#1f2933" stroke-width="4"/><circle cx="220" cy="300" r="120" fill="#dbeafe" stroke="#0057b8" stroke-width="8"/><rect x="430" y="190" width="300" height="220" rx="18" fill="#dcfce7" stroke="#15803d" stroke-width="8"/><text x="80" y="90" font-family="Arial" font-size="34" font-weight="700" fill="#1f2933">nLab DEMO SVG</text><text x="80" y="135" font-family="Arial" font-size="22" fill="#52616e">Fichier vectoriel synthetique</text></svg>'
    path.write_text(svg, encoding="utf-8")

def make_text_image(path, *, handwritten=False, angle=0, low_contrast=False, noisy=False, fmt=None):
    W,H=1200,1680
    bg=246 if low_contrast else 255
    img=Image.new("L",(W,H),bg); d=ImageDraw.Draw(img)
    title=fnt("italic" if handwritten else "bold",48)
    body=fnt("italic" if handwritten else "reg",32)
    ink=112 if low_contrast else 28
    d.text((90,80),"nLab OCR DEMO - contenu synthetique",font=title,fill=ink)
    lines=[
        "Reference DEMO-OCR-001",
        "Le renard brun saute au-dessus du chien paresseux.",
        "Lorem ipsum dolor sit amet, consectetur adipiscing elit.",
        "Montant fictif : 123,45 EUR - Date : 18/09/2026",
        "Contact fictif : demo.user@example.test",
        "Adresse fictive : 10 rue de la Demo, 00000 Ville-Test",
    ]
    y=215; rnd=random.Random(42 if handwritten else 7)
    for line in lines:
        x=90+(rnd.randint(-12,12) if handwritten else 0); yy=y+(rnd.randint(-6,6) if handwritten else 0)
        d.text((x,yy),line,font=body,fill=ink); y+=112 if handwritten else 98
    y+=25; d.rectangle((80,y,1120,y+330),outline=90 if low_contrast else 35,width=2)
    for ry in range(1,4): d.line((80,y+ry*82,1120,y+ry*82),fill=100 if low_contrast else 45,width=2)
    for cx in [300,590,860]: d.line((cx,y,cx,y+330),fill=100 if low_contrast else 45,width=2)
    for xx,txt in zip([105,330,620,890],["ID","Libelle","Valeur","Statut"]): d.text((xx,y+18),txt,font=body,fill=ink)
    for r,row in enumerate([("001","Alpha","10","OK"),("002","Beta","20","TEST"),("003","Gamma","30","OK")]):
        yy=y+82*(r+1)+18
        for xx,txt in zip([105,330,620,890],row): d.text((xx,yy),txt,font=body,fill=ink)
    if noisy:
        noise=Image.effect_noise((W,H),22).convert("L")
        img=Image.blend(img,noise,0.07).filter(ImageFilter.GaussianBlur(0.45))
        dd=ImageDraw.Draw(img); rr=random.Random(3)
        for _ in range(400):
            x,y0=rr.randrange(W),rr.randrange(H); dd.point((x,y0),fill=rr.choice([70,110,160,200]))
    if angle: img=img.rotate(angle,resample=Image.Resampling.BICUBIC,expand=True,fillcolor=255)
    ext=(fmt or path.suffix.lower().lstrip('.')).lower()
    if ext in ("jpg","jpeg"): img.save(path,"JPEG",quality=78,optimize=True)
    else: img.save(path)

def make_images(root):
    out=root/"demo-input/Images"; img=Image.new("RGB",(1000,680),"white"); d=ImageDraw.Draw(img)
    d.rectangle((35,35,965,645),outline="black",width=5); d.text((70,70),"nLab PDF Studio - IMAGE DEMO",font=fnt("bold",44),fill="black")
    d.text((70,140),"Rotation / recadrage / conversion / annotation",font=fnt("reg",28),fill="#334155")
    d.rectangle((140,250,420,520),outline="#0057b8",width=10); d.ellipse((585,245,855,515),outline="#15803d",width=10); d.line((80,590,920,590),fill="#b91c1c",width=8)
    img.save(out/"demo-image-color.png",optimize=True); img.save(out/"demo-image-color.jpg",quality=82,optimize=True); img.save(out/"demo-image-color.webp",quality=82); img.save(out/"demo-image-color.bmp")
    f2=img.copy(); ImageDraw.Draw(f2).text((70,195),"FRAME 2",font=fnt("bold",34),fill="#7c3aed"); img.save(out/"demo-image-anime.gif",save_all=True,append_images=[f2],duration=650,loop=0)
    save_svg(out/"demo-vector.svg")

def make_ocr(root):
    out=root/"demo-input/OCR"
    make_text_image(out/"ocr-imprime-propre.png")
    make_text_image(out/"ocr-imprime-faible-contraste.jpg",low_contrast=True,fmt="jpg")
    make_text_image(out/"ocr-scan-incline-bruite.jpg",angle=2.4,noisy=True,fmt="jpg")
    make_text_image(out/"ocr-pseudo-manuscrit.png",handwritten=True)
    make_text_image(out/"ocr-pseudo-manuscrit-incline.jpg",handwritten=True,angle=-3.0,noisy=True,fmt="jpg")

def make_pdfs(root):
    out=root/"demo-input/PDF"; W,H=A4
    c=canvas.Canvas(str(out/"pdf-texte-actif.pdf"),pagesize=A4)
    for pg in range(1,4):
        c.setFont("Helvetica-Bold",18); c.drawString(55,H-65,f"nLab PDF Studio - texte natif - page {pg}")
        c.setFont("Helvetica",11); c.drawString(55,H-95,"Texte natif selectionnable : surlignage, copie, recherche et comparaison.")
        for i in range(14): c.drawString(55,H-130-i*34,f"DEMO-{pg:02d}-{i+1:02d}  Valeur synthetique {100+pg*10+i}  Statut TEST")
        c.showPage()
    c.save()
    c=canvas.Canvas(str(out/"pdf-multipage-orientations.pdf"),pagesize=A4); c.setFont("Helvetica-Bold",18); c.drawString(55,H-70,"Portrait - page 1"); c.showPage(); LW,LH=landscape(A4); c.setPageSize((LW,LH)); c.drawString(55,LH-70,"Paysage - page 2"); c.showPage(); c.setPageSize(A4); c.drawString(55,H-70,"Portrait - page 3"); c.showPage(); c.save()
    c=canvas.Canvas(str(out/"pdf-formulaire-acroform.pdf"),pagesize=A4); form=c.acroForm; c.setFont("Helvetica-Bold",16); c.drawString(55,H-65,"Formulaire PDF AcroForm - DEMO"); c.setFont("Helvetica",10)
    c.drawString(55,H-105,"Identifiant synthetique :"); form.textfield(name="identifiant_demo",x=190,y=H-118,width=260,height=22,borderWidth=1)
    c.drawString(55,H-150,"Activer option :"); form.checkbox(name="option_demo",x=190,y=H-160,size=14,buttonStyle="check")
    c.drawString(55,H-195,"Categorie :"); form.choice(name="categorie_demo",x=190,y=H-208,width=190,height=22,options=["Option A","Option B","Option C"],value="Option A")
    c.drawString(55,H-245,"Commentaire :"); form.textfield(name="commentaire_demo",x=190,y=H-310,width=300,height=70,borderWidth=1,fieldFlags="multiline"); c.save()
    for src_name,out_name in [("ocr-imprime-propre.png","pdf-scan-image-only-propre.pdf"),("ocr-scan-incline-bruite.jpg","pdf-scan-image-only-incline-bruite.pdf"),("ocr-pseudo-manuscrit.png","pdf-scan-pseudo-manuscrit.pdf")]:
        im=Image.open(root/"demo-input/OCR"/src_name).convert("RGB"); im.save(out/out_name,"PDF",resolution=145)

def make_office(root):
    out=root/"demo-input/Office"; doc=Document(); title=doc.add_heading("nLab - Document Word de demonstration",level=1); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Document entierement synthetique. Aucun nom, adresse ou identifiant reel."); doc.add_heading("1. Texte et styles",level=2)
    p=doc.add_paragraph(); r=p.add_run("Gras "); r.bold=True; r=p.add_run("Italique "); r.italic=True; r=p.add_run("Souligne"); r.underline=True
    doc.add_paragraph("Lorem ipsum dolor sit amet, consectetur adipiscing elit. Reference DEMO-DOCX-001."); doc.add_heading("2. Tableau",level=2)
    t=doc.add_table(rows=1,cols=4)
    for i,h in enumerate(["Reference","Date","Valeur","Statut"]): t.rows[0].cells[i].text=h
    for i in range(1,7):
        cells=t.add_row().cells; cells[0].text=f"DEMO-{i:03d}"; cells[1].text=f"2026-09-{i:02d}"; cells[2].text=str(i*12.5); cells[3].text=["Nouveau","En cours","Termine"][i%3]
    doc.add_heading("3. Liste",level=2)
    for v in ["Import DOCX","Conversion PDF","Extraction texte","Tests de styles"]: doc.add_paragraph(v,style="List Bullet")
    doc.save(out/"document-demo.docx")
    (out/"document-demo.rtf").write_text(r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Arial;}} \fs28 nLab RTF DEMO\par \fs22 Contenu synthetique sans donnee personnelle.\par DEMO-RTF-001\par}",encoding="latin-1")
    odt=out/"document-demo.odt"
    content = '<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:text><text:h text:outline-level="1">nLab ODT DEMO</text:h><text:p>Contenu synthetique sans donnee personnelle.</text:p><text:p>DEMO-ODT-001</text:p></office:text></office:body></office:document-content>'
    manifest = '<?xml version="1.0" encoding="UTF-8"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>'
    with zipfile.ZipFile(odt,"w") as z:
        zi=zipfile.ZipInfo("mimetype"); zi.compress_type=zipfile.ZIP_STORED; z.writestr(zi,"application/vnd.oasis.opendocument.text")
        z.writestr("content.xml",content); z.writestr("META-INF/manifest.xml",manifest)

def make_xlsx(path):
    wb=Workbook(); ws=wb.active; ws.title="Clients"
    ws.append(["Client_ID","Nom","Prenom","Email","Telephone","Ville","Statut","Montant","Date"])
    for i in range(1,31):
        ws.append([
            f"DEMO-C{i:03d}", f"NomDemo{i:03d}", f"PrenomDemo{i:03d}",
            f"client{i:03d}@example.test", f"060000{i:04d}", f"Ville-Test-{(i%5)+1}",
            ["Nouveau","Actif","Archive"][i%3], round(i*17.35,2),
            datetime(2026,9,1+((i-1)%20))
        ])
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="1F2933")
        cell.alignment=Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2,max_row=31,min_col=8,max_col=8):
        row[0].number_format='#,##0.00 "EUR"'
    for row in ws.iter_rows(min_row=2,max_row=31,min_col=9,max_col=9):
        row[0].number_format="yyyy-mm-dd"
    widths=[16,16,18,28,16,18,14,14,14]
    for idx,w in enumerate(widths,1):
        ws.column_dimensions[chr(64+idx)].width=w
    ws.freeze_panes="A2"
    tab=Table(displayName="ClientsDemo",ref="A1:I31")
    tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showFirstColumn=False,showLastColumn=False,showColumnStripes=False)
    ws.add_table(tab)
    dv=DataValidation(type="list",formula1='"Nouveau,Actif,Archive"',allow_blank=True)
    ws.add_data_validation(dv); dv.add("G2:G31")
    rs=wb.create_sheet("Resume")
    rs.append(["Indicateur","Valeur"]); rs.append(["Nombre de clients",30])
    rs.append(["Somme montants",sum(round(i*17.35,2) for i in range(1,31))])
    rs.append(["Date min",datetime(2026,9,1)]); rs.append(["Date max",datetime(2026,9,20)])
    for cell in rs[1]:
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="0057B8")
    rs["B3"].number_format='#,##0.00 "EUR"'
    rs["B4"].number_format="yyyy-mm-dd"; rs["B5"].number_format="yyyy-mm-dd"
    rs.column_dimensions["A"].width=24; rs.column_dimensions["B"].width=18
    wb.save(path)

def make_data(root):
    out=root/"demo-input/Data"; make_xlsx(out/"donnees-synthetiques.xlsx")
    rec=[]
    for i in range(1,16): rec.append({"client_id":f"DEMO-C{i:03d}","nom":f"NomDemo{i:03d}","prenom":f"PrenomDemo{i:03d}","email":f"client{i:03d}@example.test","telephone":f"060000{i:04d}","adresse":f"{i} rue de la Demo","code_postal":"00000","ville":f"Ville-Test-{(i%5)+1}","montant":round(i*17.35,2),"statut":["Nouveau","Actif","Archive"][i%3]})
    with open(out/"clients-demo.csv","w",encoding="utf-8",newline="") as f: w=csv.DictWriter(f,fieldnames=list(rec[0])); w.writeheader(); w.writerows(rec)
    (out/"clients-demo.json").write_text(json.dumps(rec,ensure_ascii=False,indent=2),encoding="utf-8"); (out/"clients-demo.ndjson").write_text("\n".join(json.dumps(x,ensure_ascii=False) for x in rec)+"\n",encoding="utf-8")
    (out/"clients-demo.yaml").write_text("dataset: nLab-demo-v2\nrecords:\n"+"".join(f"  - client_id: {r['client_id']}\n    nom: {r['nom']}\n    statut: {r['statut']}\n" for r in rec[:6]),encoding="utf-8")
    xml=['<?xml version="1.0" encoding="UTF-8"?>','<clients dataset="nLab-demo-v2">']
    for r in rec[:8]: xml += [f'  <client id="{r["client_id"]}">',f'    <nom>{r["nom"]}</nom>',f'    <statut>{r["statut"]}</statut>','  </client>']
    xml.append('</clients>'); (out/"clients-demo.xml").write_text("\n".join(xml),encoding="utf-8"); (out/"notes-demo.txt").write_text("nLab DEMO v2\nAucune donnee personnelle reelle.\nReferences : DEMO-TXT-001, DEMO-TXT-002.\n",encoding="utf-8")

def make_compare(root):
    a=root/"demo-input/Compare/tree-A"; b=root/"demo-input/Compare/tree-B"
    (a/"docs/commun.txt").write_text("Meme contenu dans A et B.\n",encoding="utf-8"); (b/"docs/commun.txt").write_text("Meme contenu dans A et B.\n",encoding="utf-8")
    (a/"docs/modifie.txt").write_text("Version A - contenu initial.\n",encoding="utf-8"); (b/"docs/modifie.txt").write_text("Version B - contenu modifie.\n",encoding="utf-8")
    (a/"docs/seulement-A.txt").write_text("Present uniquement dans A.\n",encoding="utf-8"); (b/"docs/seulement-B.txt").write_text("Present uniquement dans B.\n",encoding="utf-8")
    (a/"data/liste.csv").write_text("id,valeur\n1,10\n2,20\n",encoding="utf-8"); (b/"data/liste.csv").write_text("id,valeur\n1,10\n2,21\n",encoding="utf-8")

def make_nested_zip(root):
    out=root/"demo-input/Zip/archive-imbriquee-demo.zip"; tiny=io.BytesIO(); Image.new("RGB",(320,180),"white").save(tiny,"PNG")
    with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
        z.writestr("sous-dossier/readme.txt","Archive imbriquee de demonstration.\n"); z.writestr("sous-dossier/data.json",json.dumps({"demo":True,"id":"DEMO-ZIP-001"},indent=2)); z.writestr("racine.csv","cle,valeur\nalpha,1\nbeta,2\n"); z.writestr("sous-dossier/image-demo.png",tiny.getvalue())

def sha256(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for ch in iter(lambda:f.read(1024*1024),b""): h.update(ch)
    return h.hexdigest()

def make_docs(root):
    (root/"demo-output/README.txt").write_text("Dossier de sortie logique de demonstration. Sur le web, les resultats utilisent les Telechargements du navigateur. Le navigateur ne peut pas imposer silencieusement un dossier local precis.\n",encoding="utf-8")
    (root/"README-DEMO.md").write_text("# nLab DEMO CORPUS v2\n\nJeu de fichiers 100 % synthetiques destine aux tests de nLab PDF Studio et des autres applications nLab.\n\n- demo-input/ : corpus a charger ;\n- demo-output/ : destination logique ;\n- docs/expected-results.json : cas de test.\n",encoding="utf-8")
    expected={"schema":"nlab-demo-corpus/v2","version":"2.0","privacy":"100% synthetic","defaultOutputMode":"browser-downloads","tests":[{"file":"demo-input/PDF/pdf-texte-actif.pdf","expect":"native text selectable"},{"file":"demo-input/PDF/pdf-formulaire-acroform.pdf","expect":"detect/fill/flatten form fields"},{"file":"demo-input/PDF/pdf-scan-image-only-incline-bruite.pdf","expect":"OCR challenge; no native text"},{"file":"demo-input/OCR/ocr-pseudo-manuscrit.png","expect":"handwriting-like OCR challenge"},{"file":"demo-input/Images/demo-image-color.webp","expect":"open image; rotate; convert to PDF"},{"file":"demo-input/Zip/archive-imbriquee-demo.zip","expect":"nested ZIP expansion test"}]}
    (root/"docs/expected-results.json").write_text(json.dumps(expected,ensure_ascii=False,indent=2),encoding="utf-8")

def build(output_dir):
    output_dir=Path(output_dir); output_dir.mkdir(parents=True,exist_ok=True); work=output_dir/"_build-demo-v2"
    if work.exists(): shutil.rmtree(work)
    ensure_dirs(work); make_images(work); make_ocr(work); make_pdfs(work); make_office(work); make_data(work); make_compare(work); make_nested_zip(work); make_docs(work)
    files=[p for p in sorted(work.rglob("*")) if p.is_file()]
    catalog=[]
    for p in files:
        rel=p.relative_to(work).as_posix(); cat=(rel.split('/')[1] if rel.startswith('demo-input/') else ('Output' if rel.startswith('demo-output/') else 'Documentation'))
        catalog.append({"path":rel,"category":cat,"size":p.stat().st_size,"sha256":sha256(p)})
    archive=output_dir/"nLab-DEMO-CORPUS-v2.zip"
    if archive.exists(): archive.unlink()
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED,compresslevel=7) as z:
        for p in files: z.write(p,p.relative_to(work).as_posix())
    manifest={"schema":"nlab-demo-manifest/v2","version":"2.0","label":"nLab DEMO CORPUS v2","privacy":"synthetic-only","archive":"../../Library/demo/nLab-DEMO-CORPUS-v2.zip","loadRoot":"demo-input/","pdfStudioExtensions":["pdf","png","jpg","jpeg","webp","gif","bmp","docx","zip"],"defaultOutputMode":"download","fileCount":len(files),"catalog":"../../Library/demo/demo-catalog-v2.json"}
    (output_dir/"demo-manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    (output_dir/"demo-catalog-v2.json").write_text(json.dumps({"version":"2.0","files":catalog},ensure_ascii=False,indent=2),encoding="utf-8")
    shutil.rmtree(work)
    print(f"Generated {archive} ({archive.stat().st_size} bytes, {len(files)} files)")

if __name__ == "__main__":
    ap=argparse.ArgumentParser(); ap.add_argument("--output",default="Library/demo"); args=ap.parse_args(); build(args.output)
