#!/usr/bin/env python3
"""Enrich the nLab public synthetic corpus for nLab Studios.
Runs after generate_demo.py. All generated identities, addresses, emails, phones and IDs are TEST data.
"""
from pathlib import Path
import argparse, csv, hashlib, json, math, shutil, zipfile
from PIL import Image, ImageDraw, ImageFont
from reportlab.graphics.barcode import createBarcodeDrawing
from reportlab.graphics import renderSVG
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

PUBLIC_BASE="https://nepheris.github.io/nLab-web/Library/demo/files/demo-input"
REG=["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf","/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"]
BOLD=["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf","/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"]

def font(names,size):
    for p in names:
        if Path(p).exists():
            return ImageFont.truetype(p,size=size)
    return ImageFont.load_default()

def fnt(bold,size): return font(BOLD if bold else REG,size)

def sha256(path):
    h=hashlib.sha256()
    with open(path,"rb") as fh:
        for chunk in iter(lambda:fh.read(1024*1024),b""): h.update(chunk)
    return h.hexdigest()

def mkdirs(root):
    for d in ["Images","Data/images","Code","QR-Barcode","File/lot-A","File/lot-B/sous-dossier","Compare/tree-A/docs","Compare/tree-B/docs"]:
        (root/"files/demo-input"/d).mkdir(parents=True,exist_ok=True)
    (root/"manifests").mkdir(parents=True,exist_ok=True)
    (root/"packs").mkdir(parents=True,exist_ok=True)

def make_calibration(root):
    out=root/"files/demo-input/Images"; ppm=4.0; W,H=1400,900; x0,y0=120,780
    img=Image.new("RGB",(W,H),"white"); d=ImageDraw.Draw(img)
    d.text((55,32),"nLab Image Studio — DEMO TEST CALIBRATION",font=fnt(True,34),fill="#1f2933")
    d.text((55,78),"Règle X 300 mm · Règle Y 150 mm · échelle attendue 4 px/mm",font=fnt(False,22),fill="#52616e")
    d.line((x0,y0,x0+1200,y0),fill="#111827",width=3)
    for mm in range(301):
        x=x0+round(mm*ppm); tick=28 if mm%50==0 else 18 if mm%10==0 else 8
        d.line((x,y0,x,y0-tick),fill="#111827",width=2 if mm%10==0 else 1)
        if mm%50==0: d.text((x-12,y0-60),str(mm),font=fnt(False,16),fill="#111827")
    d.line((x0,y0,x0,y0-600),fill="#111827",width=3)
    for mm in range(151):
        y=y0-round(mm*ppm); tick=28 if mm%50==0 else 18 if mm%10==0 else 8
        d.line((x0,y,x0+tick,y),fill="#111827",width=2 if mm%10==0 else 1)
        if mm%50==0: d.text((x0+34,y-10),str(mm),font=fnt(False,16),fill="#111827")
    ox,oy=x0+300,y0-460; ow,oh=480,320
    d.rectangle((ox,oy,ox+ow,oy+oh),outline="#0057b8",width=5)
    d.text((ox+22,oy+24),"OBJET TEST",font=fnt(True,24),fill="#0057b8")
    d.text((ox+22,oy+60),"120 × 80 mm",font=fnt(False,22),fill="#0057b8")
    path=out/"calibration-regles-300x150.png"; img.save(path,optimize=True)
    (out/"calibration-regles-300x150.json").write_text(json.dumps({
        "schema":"nlab-image-calibration-demo/v1","label":"NLAB DEMO TEST — règles 300 × 150 mm",
        "pixels_per_mm_expected":ppm,"dpi_expected":ppm*25.4,
        "x_reference_mm":300,"y_reference_mm":150,
        "known_object_mm":{"width":120,"height":80},"privacy":"synthetic-only"
    },ensure_ascii=False,indent=2),encoding="utf-8")

def make_code(root):
    out=root/"files/demo-input/Code"
    (out/"demo-script.js").write_text(
        "const rows = [{id:'TEST-001', value:12.5}, {id:'TEST-002', value:18.75}];\n"
        "console.table(rows);\nconsole.log(rows.reduce((s,r)=>s+r.value,0));\n",encoding="utf-8")
    (out/"demo-structure.json").write_text(json.dumps({
        "dataset":"NLAB_DEMO_TEST","synthetic":True,
        "records":[{"id":"TEST-001","nom":"Alpha","tags":["demo","json"]},{"id":"TEST-002","nom":"Bêta","tags":["unicode","test"]}]
    },ensure_ascii=False,indent=2),encoding="utf-8")
    (out/"demo-markdown.md").write_text("# nLab Code Studio — DEMO TEST\n\n- JavaScript\n- JSON\n- Markdown\n- YAML\n- XML\n",encoding="utf-8")
    (out/"demo-config.yaml").write_text("dataset: NLAB_DEMO_TEST\nsynthetic: true\nfeatures:\n  - code\n  - json\n",encoding="utf-8")
    (out/"demo-record.xml").write_text('<?xml version="1.0" encoding="UTF-8"?>\n<demo synthetic="true"><id>TEST-001</id><label>nLab Demo</label></demo>\n',encoding="utf-8")

def make_codes(root):
    out=root/"files/demo-input/QR-Barcode"
    items=[
        ("qr-nlab","QR","https://nepheris.github.io/nLab-web/","nLab Web"),
        ("qr-youtube","QR","https://www.youtube.com/","YouTube — URL publique"),
        ("qr-lemonde","QR","https://www.lemonde.fr/","Le Monde — URL publique"),
        ("qr-google","QR","https://www.google.com/","Google — URL publique"),
        ("datamatrix-test","ECC200DataMatrix","NLAB-DEMO-DATAMATRIX-001","Data Matrix synthétique"),
        ("code128-test","Code128","NLAB-DEMO-CODE128-001","Code 128 synthétique")
    ]
    meta=[]
    for fid,kind,value,label in items:
        if kind=="Code128":
            drawing=createBarcodeDrawing(kind,value=value,width=440,height=150,humanReadable=True)
        else:
            drawing=createBarcodeDrawing(kind,value=value,width=260,height=260)
        renderSVG.drawToFile(drawing,str(out/f"{fid}.svg"))
        meta.append({"id":fid,"symbology":"DATAMATRIX" if kind=="ECC200DataMatrix" else kind,"value":value,"label":label,"file":f"{fid}.svg"})
    (out/"payloads.json").write_text(json.dumps({"schema":"nlab-code-demo/v1","privacy":"synthetic-or-public-url-only","items":meta},ensure_ascii=False,indent=2),encoding="utf-8")

def make_file_samples(root):
    a=root/"files/demo-input/File/lot-A"; b=root/"files/demo-input/File/lot-B/sous-dossier"
    samples={
        a/"IMG_0001 photo ROUGE.txt":"DEMO TEST fichier 1\n",
        a/"IMG_0002 photo BLEUE.txt":"DEMO TEST fichier 2\n",
        a/"facture test 0003.txt":"DEMO TEST facture fictive\n",
        a/"rapport.v1.final.TEST.txt":"DEMO TEST nom avec points\n",
        b/"2026_09_22-client_demo-A.txt":"DEMO TEST classement\n",
        b/"Nom Avec Espaces et accents É.txt":"DEMO TEST unicode\n",
        b/"data_TEST_001.json":json.dumps({"demo":True,"id":"FILE-001"},indent=2),
        b/"liste_TEST.csv":"id,label\n1,Alpha\n2,Beta\n"
    }
    for path,content in samples.items(): path.write_text(content,encoding="utf-8")
    ca=root/"files/demo-input/Compare/tree-A/docs"; cb=root/"files/demo-input/Compare/tree-B/docs"
    (ca/"commun.txt").write_text("Même contenu TEST.\n",encoding="utf-8")
    (cb/"commun.txt").write_text("Même contenu TEST.\n",encoding="utf-8")
    (ca/"modifie.txt").write_text("Version A — DEMO TEST.\n",encoding="utf-8")
    (cb/"modifie.txt").write_text("Version B — DEMO TEST modifiée.\n",encoding="utf-8")

def luhn_ok(num):
    total=0
    for i,ch in enumerate(num):
        n=int(ch)
        if i%2==0:
            n*=2
            if n>9:n-=9
        total+=n
    return total%10==0

def invalid_siret(i):
    n=f"1234567890{i:04d}"[-14:]
    if luhn_ok(n): n=n[:-1]+str((int(n[-1])+1)%10)
    return n

def make_dataset(root):
    out=root/"files/demo-input/Data"; images=out/"images"
    palette=[("Rouge test","#D32F2F"),("Bleu test","#1976D2"),("Vert test","#388E3C"),("Orange test","#F57C00"),("Violet test","#7B1FA2"),("Turquoise test","#00838F"),("Gris test","#616161"),("Jaune test","#FBC02D")]
    for name,hexv in palette:
        im=Image.new("RGB",(180,90),hexv); d=ImageDraw.Draw(im)
        d.rectangle((0,0,179,89),outline="white",width=3)
        d.text((8,32),name,font=fnt(True,16),fill="black" if "Jaune" in name else "white")
        im.save(images/(name.lower().replace(" ","-")+".png"))
    public_urls=["https://www.google.com/","https://www.youtube.com/","https://www.lemonde.fr/","https://fr.wikipedia.org/","https://www.insee.fr/","https://www.data.gouv.fr/"]
    first=["Alice-Test","Benoît-Test","Chloé-Test","David-Test","Élodie-Test","Farid-Test"]
    last=["Dataset","Démonstration","Synthétique","Exemple","Validation","nLab"]
    rows=[]
    for i in range(1,41):
        cname,chex=palette[(i-1)%len(palette)]; imgname=cname.lower().replace(" ","-")+".png"
        rows.append({
            "test_id":f"NLAB-TEST-{i:04d}",
            "duplicate_key":f"GROUP-{((i-1)%6)+1}",
            "first_name_test":first[(i-1)%len(first)],
            "last_name_test":last[(i*2)%len(last)],
            "display_name_test":f"{first[(i-1)%len(first)]} {last[(i*2)%len(last)]}",
            "color_name":cname,"color_hex":chex,
            "address_test":f"{i} rue du Dataset de Test, 00000 Ville-Démo",
            "email_safe":f"utilisateur.test.{i:03d}@example.test",
            "email_plus_safe":f"dataset+ligne{i:03d}@example.test",
            "email_upper_safe":f"DEMO.TEST.{i:03d}@EXAMPLE.TEST",
            "phone_test_invalid":f"+33 0 00 00 00 {i:02d}",
            "siret_test_invalid":invalid_siret(i),
            "category_discrete":["A","B","C","D"][i%4],
            "status_discrete":["NOUVEAU","EN_COURS","TERMINE"][i%3],
            "integer_value":i*3,
            "continuous_value":round(math.sin(i/4)*100+i/7,6),
            "percentage":round(((i*7.3)%100)/100,4),
            "money_eur":round(i*17.35+0.005*(i%3),3),
            "rounding_case":["0.1 + 0.2","1.005 → 1.01 ?","2.675 → 2.68 ?"][i%3],
            "accounting_signed":round(((-1)**i)*i*23.47,2),
            "date_iso":f"2026-09-{((i-1)%28)+1:02d}",
            "datetime_iso":f"2026-09-{((i-1)%28)+1:02d}T{(8+i)%24:02d}:{(i*7)%60:02d}:00+02:00",
            "boolean_value":i%2==0,
            "nullable_value":None if i%7==0 else f"VAL-{i:03d}",
            "tags_list":json.dumps(["demo","test",f"groupe-{i%4}"],ensure_ascii=False),
            "public_url":public_urls[(i-1)%len(public_urls)],
            "image_url":f"{PUBLIC_BASE}/Data/images/{imgname}",
            "lorem_text":"Lorem ipsum dolor sit amet, consectetur adipiscing elit. Dataset de test nLab."
        })
    fields=list(rows[0])
    with open(out/"dataset-validation-complet.csv","w",encoding="utf-8",newline="") as fh:
        w=csv.DictWriter(fh,fieldnames=fields); w.writeheader(); w.writerows(rows)
    (out/"dataset-validation-complet.json").write_text(json.dumps({"schema":"nlab-demo-dataset/v1","label":"NLAB DATASET DE TEST — SYNTHETIQUE","records":rows},ensure_ascii=False,indent=2),encoding="utf-8")
    wb=Workbook(); ws=wb.active; ws.title="Dataset_Test"; xfields=fields+["image_preview"]; ws.append(xfields)
    for c in ws[1]:
        c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="0057B8"); c.alignment=Alignment(horizontal="center")
    for rix,row in enumerate(rows,2):
        ws.append([row[k] for k in fields]+[""])
        for key in ["public_url","image_url"]:
            c=ws.cell(rix,fields.index(key)+1); c.hyperlink=row[key]; c.style="Hyperlink"
        ws.cell(rix,fields.index("percentage")+1).number_format="0.00%"
        ws.cell(rix,fields.index("money_eur")+1).number_format='#,##0.00 "€"'
        ws.cell(rix,fields.index("accounting_signed")+1).number_format='#,##0.00 "€";[Red]-#,##0.00 "€"'
        ws.cell(rix,fields.index("continuous_value")+1).number_format="0.000000"
        pic=XLImage(images/(row["color_name"].lower().replace(" ","-")+".png")); pic.width=72; pic.height=36
        ws.add_image(pic,f"{get_column_letter(len(xfields))}{rix}"); ws.row_dimensions[rix].height=32
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for col in range(1,len(xfields)+1): ws.column_dimensions[get_column_letter(col)].width=20
    meta=wb.create_sheet("Dictionnaire"); meta.append(["colonne","type attendu","propriété de test"])
    dictionary=[
        ("test_id","string","clé unique"),("duplicate_key","string","clé volontairement non unique"),
        ("color_name / color_hex","catégorie / couleur","rendu couleur"),("email_*","email","formats sûrs en .test"),
        ("phone_test_invalid","texte","numéro volontairement non attribué"),("siret_test_invalid","14 chiffres","forme SIRET mais contrôle Luhn invalide"),
        ("continuous_value","float","variable continue"),("category_discrete","catégorie","variable discrète"),
        ("percentage","float","format pourcentage"),("money_eur","monétaire","format comptable et arrondis"),
        ("public_url","URL","liens publics cliquables"),("image_url","URL image","images dans un tableau"),
        ("nullable_value","nullable","valeurs manquantes"),("tags_list","liste sérialisée","multi-valeurs")
    ]
    for r in dictionary: meta.append(r)
    wb.save(out/"dataset-validation-complet.xlsx")

def collect(root):
    files=[]
    for base in ["files/demo-input","files/demo-output","docs"]:
        p=root/base
        if p.exists():
            files += [x for x in sorted(p.rglob("*")) if x.is_file()]
    return files

def rebuild_outputs(root):
    files=collect(root)
    catalog=[]
    for p in files:
        rel=p.relative_to(root).as_posix()
        cat=rel.split("/")[2] if rel.startswith("files/demo-input/") and len(rel.split("/"))>2 else ("Output" if rel.startswith("files/demo-output/") else "Documentation")
        catalog.append({"path":rel,"category":cat,"size":p.stat().st_size,"sha256":sha256(p)})
    archive=root/"nLab-DEMO-CORPUS-v3.zip"
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED,compresslevel=7) as z:
        for p in files:z.write(p,p.relative_to(root).as_posix())
    specs={
        "pdf":{"label":"PDF Studio","prefixes":["files/demo-input/PDF/","files/demo-input/OCR/","files/demo-input/Zip/"]},
        "image":{"label":"Image Studio","prefixes":["files/demo-input/Images/"]},
        "ocr":{"label":"OCR Studio","prefixes":["files/demo-input/OCR/"],"also":["files/demo-input/PDF/pdf-scan-image-only-propre.pdf","files/demo-input/PDF/pdf-scan-image-only-incline-bruite.pdf"]},
        "code-json":{"label":"Code + JSON Studio","prefixes":["files/demo-input/Code/"],"also":["files/demo-input/Data/dataset-validation-complet.json"]},
        "data":{"label":"Data Studio","prefixes":["files/demo-input/Data/"]},
        "qr-barcode":{"label":"QR & Barcode Studio","prefixes":["files/demo-input/QR-Barcode/"]},
        "file":{"label":"File Studio","prefixes":["files/demo-input/File/","files/demo-input/Compare/","files/demo-input/Zip/"]}
    }
    idx={"schema":"nlab-studio-demo-index/v1","version":"3.0","privacy":"synthetic-only","studios":[]}
    for sid,spec in specs.items():
        chosen=[p for p in files if any(p.relative_to(root).as_posix().startswith(pref) for pref in spec.get("prefixes",[])) or p.relative_to(root).as_posix() in spec.get("also",[])]
        pack=root/"packs"/f"nLab-DEMO-{sid}-v3.zip"
        with zipfile.ZipFile(pack,"w",zipfile.ZIP_DEFLATED,compresslevel=7) as z:
            for p in chosen:z.write(p,p.relative_to(root).as_posix())
        sm={"schema":"nlab-studio-demo-manifest/v1","version":"3.0","studio":sid,"label":spec["label"],"privacy":"synthetic-only","pack":f"../packs/{pack.name}","files":[{"path":"../"+p.relative_to(root).as_posix(),"size":p.stat().st_size,"sha256":sha256(p)} for p in chosen]}
        if sid=="file": sm["demoNames"]=["IMG_0001 photo ROUGE.txt","IMG_0002 photo BLEUE.txt","facture test 0003.txt","rapport.v1.final.TEST.txt","Nom Avec Espaces et accents É.txt"]
        (root/"manifests"/f"{sid}.json").write_text(json.dumps(sm,ensure_ascii=False,indent=2),encoding="utf-8")
        idx["studios"].append({"id":sid,"label":spec["label"],"manifest":f"./{sid}.json","pack":f"../packs/{pack.name}","fileCount":len(chosen)})
    (root/"manifests/index.json").write_text(json.dumps(idx,ensure_ascii=False,indent=2),encoding="utf-8")
    gallery=[]
    for item in catalog:
        ext=Path(item["path"]).suffix.lower()
        preview="image" if ext in [".png",".jpg",".jpeg",".webp",".gif",".svg"] else "json" if ext==".json" else "csv" if ext==".csv" else "pdf" if ext==".pdf" else "text" if ext in [".txt",".md",".yaml",".yml",".xml",".js",".css",".html",".py",".sql"] else "download"
        gallery.append({**item,"href":"./"+item["path"],"preview":preview})
    (root/"demo-gallery-v3.json").write_text(json.dumps({"schema":"nlab-demo-gallery/v1","version":"3.0","files":gallery},ensure_ascii=False,indent=2),encoding="utf-8")
    base_manifest={}
    base_path=root/"demo-manifest.json"
    if base_path.exists():
        base_manifest=json.loads(base_path.read_text(encoding="utf-8"))
    base_manifest["studioDemo"]={
        "version":"3.0",
        "archive":"../../Library/demo/nLab-DEMO-CORPUS-v3.zip",
        "catalog":"../../Library/demo/demo-catalog-v3.json",
        "gallery":"../../Library/demo/demo-gallery-v3.json",
        "manifests":"../../Library/demo/manifests/index.json",
        "privacy":"synthetic-only"
    }
    base_manifest["studioManifests"]="../../Library/demo/manifests/index.json"
    base_manifest["galleryV3"]="../../Library/demo/demo-gallery-v3.json"
    (root/"demo-manifest.json").write_text(json.dumps(base_manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    (root/"demo-catalog-v3.json").write_text(json.dumps({"version":"3.0","files":catalog},ensure_ascii=False,indent=2),encoding="utf-8")

def build(output):
    root=Path(output); mkdirs(root); make_calibration(root); make_code(root); make_codes(root); make_file_samples(root); make_dataset(root); rebuild_outputs(root)
    print("nLab DEMO CORPUS v3 enriched:",root)

if __name__=="__main__":
    ap=argparse.ArgumentParser(); ap.add_argument("--output",default="Library/demo"); args=ap.parse_args(); build(args.output)
